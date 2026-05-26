import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def generate_calificaciones_excel(materia_id, datos):
    # Función heredada simplificada para compatibilidad
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calificaciones"
    
    ws['A1'] = f"Reporte Final de Calificaciones - Materia ID: {materia_id}"
    ws['A1'].font = Font(size=14, bold=True)
    ws.merge_cells('A1:D1')
    
    headers = ["Matrícula", "Nombre del Alumno", "Asistencia (%)", "Calificación Final"]
    ws.append([])
    ws.append(headers)
    
    for col in range(1, 5):
        cell = ws.cell(row=3, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 18
    
    for alumno in datos:
        ws.append([
            alumno.get('matricula', 'N/A'),
            alumno.get('nombre', 'Desconocido'),
            alumno.get('asistencia', 0),
            alumno.get('calificacion', 0.0)
        ])
        
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def generate_asistencias_excel(materia_id, datos):
    # Función heredada simplificada para compatibilidad
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Asistencias"
    
    ws['A1'] = f"Reporte de Asistencias - Materia ID: {materia_id}"
    ws['A1'].font = Font(size=14, bold=True)
    ws.merge_cells('A1:E1')

    headers = ["Matrícula", "Nombre del Alumno", "Presentes", "Retardos", "Faltas"]
    ws.append([]) 
    ws.append(headers)

    for col in range(1, 6):
        cell = ws.cell(row=3, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")

    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12

    for alumno in datos:
        ws.append([
            alumno.get('matricula', 'N/A'),
            alumno.get('nombre', 'Desconocido'),
            alumno.get('presentes', 0),
            alumno.get('retardos', 0),
            alumno.get('faltas', 0)
        ])
    
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def generate_rendimiento_excel(materia_id, resumen, alumnos):
    # Función heredada simplificada para compatibilidad
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rendimiento"

    ws['A1'] = f"Reporte de Rendimiento - Materia ID: {materia_id}"
    ws['A1'].font = Font(size=14, bold=True)
    ws.merge_cells('A1:F1')

    ws.append([])
    ws.append(["Periodo", resumen.get('periodo_id', 'N/A')])
    ws.append(["Materia", resumen.get('materia_nombre', 'N/A')])
    ws.append(["Promedio de Grupo", resumen.get('promedio_grupo', 0.0)])
    ws.append(["Calificación Máxima", resumen.get('calificacion_maxima', 0.0)])
    ws.append(["Calificación Mínima", resumen.get('calificacion_minima', 0.0)])
    ws.append(["Tasa de Asistencia (%)", resumen.get('tasa_asistencia', 0.0)])
    ws.append(["Total de Alumnos", resumen.get('total_alumnos', 0)])

    ws.append([])
    headers = ["Matrícula", "Nombre del Alumno", "Asistencia (%)", "Calificación Final", "Presentes", "Retardos", "Faltas"]
    ws.append(headers)

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=ws.max_row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

    for alumno in alumnos:
        ws.append([
            alumno.get('matricula', 'N/A'),
            alumno.get('nombre', 'Desconocido'),
            alumno.get('asistencia', 0.0),
            alumno.get('calificacion', 0.0),
            alumno.get('presentes', 0),
            alumno.get('retardos', 0),
            alumno.get('faltas', 0),
        ])

    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


# =============================================================================
# NUEVO GENERADOR CONSOLIDADO CON DOS PESTAÑAS Y ESTILOS BUAP DE ALTA GAMA
# =============================================================================

MESES = {
    1: "ene", 2: "feb", 3: "mar", 4: "abr", 5: "may", 6: "jun",
    7: "jul", 8: "ago", 9: "sep", 10: "oct", 11: "nov", 12: "dic"
}

def _format_fecha_buap(fecha_str):
    """Convierte '2026-02-02' o '2026-02-02T00:00:00' a formato '02-feb-26'."""
    try:
        from datetime import datetime
        clean_date = fecha_str.split("T")[0]
        dt = datetime.strptime(clean_date, "%Y-%m-%d")
        mes = MESES[dt.month]
        return f"{dt.day:02d}-{mes}-{dt.strftime('%y')}"
    except Exception:
        return fecha_str


def generate_consolidated_excel(materia_id, datos_calificaciones, datos_asistencias, periodo_nombre, docente_nombre):
    """
    Genera un libro de Excel consolidado de alta gama con dos pestañas:
    1. Calificaciones: Detalle por actividad (escala 0-1) y ponderación + subtotal amarillo y nota final.
    2. Asistencias: Desglose cronológico por fecha de asistencia (1 = presente, 0.5 = retardo, 0 = falta).
    """
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------------------
    # HOJA 1: CALIFICACIONES
    # -------------------------------------------------------------------------
    ws_calif = wb.active
    ws_calif.title = "Calificaciones"
    ws_calif.views.sheetView[0].showGridLines = True
    
    # Estilos de borde y fuentes
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )
    bold_font = Font(name="Calibri", size=10, bold=True)
    regular_font = Font(name="Calibri", size=10)
    
    fill_blue_light = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    fill_yellow = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    
    # 1. Cabecera (Filas 1 a 4)
    # A1:B4 para logo
    ws_calif.merge_cells("A1:B4")
    cell_logo = ws_calif["A1"]
    cell_logo.value = "BUAP"
    cell_logo.alignment = Alignment(horizontal="center", vertical="center")
    cell_logo.font = Font(name="Calibri", size=16, bold=True, italic=True, color="1F4E79")
    cell_logo.fill = fill_blue_light
    
    # Recolectar ponderaciones y actividades
    ponderaciones = datos_calificaciones.get("ponderaciones", [])
    materia_nombre = datos_calificaciones.get("materia_nombre", "DESARROLLO DE APLICACIONES WEB").upper()
    periodo_str = (periodo_nombre or "PRIMAVERA 2026").upper()
    docente_str = (docente_nombre or "M.C. LUIS YAEL MÉNDEZ SÁNCHEZ").upper()
    
    # Calcular cantidad total de columnas dinámicas para saber el fin de la cabecera
    # Columnas base: Matrícula (A), Nombre (B) -> 2
    # Por cada ponderación: actividades + 1 (Total categoría)
    num_actividades_total = 0
    for p in ponderaciones:
        num_actividades_total += len(p.get("actividades", [])) + 1
        
    last_col_idx = 2 + num_actividades_total + 1 # +1 para la Calificación Final
    last_col_letter = get_column_letter(last_col_idx)
    header_end_letter = get_column_letter(last_col_idx - 1)
    
    # Escribir textos de cabecera combinada
    cabecera_textos = [
        "BENEMÉRITA UNIVERSIDAD AUTÓNOMA DE PUEBLA",
        f"{materia_nombre} - {periodo_str}",
        "GRUPO 03:00 P.M.",
        docente_str
    ]
    
    for idx, text in enumerate(cabecera_textos, start=1):
        range_str = f"C{idx}:{header_end_letter}{idx}"
        ws_calif.merge_cells(range_str)
        cell = ws_calif[f"C{idx}"]
        cell.value = text
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = fill_blue_light
        
    # Aplicar fondo azul y bordes a toda la zona de cabecera A1:LastCol4
    for r in range(1, 5):
        for c in range(1, last_col_idx + 1):
            cell = ws_calif.cell(row=r, column=c)
            cell.fill = fill_blue_light
            cell.border = thin_border
            
    # Combinar celda de Calificación Final en cabecera
    final_header_cell = ws_calif[f"{last_col_letter}1"]
    final_header_cell.value = "FINAL"
    final_header_cell.font = Font(name="Calibri", size=10, bold=True)
    final_header_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_calif.merge_cells(f"{last_col_letter}1:{last_col_letter}4")
    
    # 2. Encabezados de la Tabla (Filas 5 a 7)
    # A5:A7 para Matrícula
    ws_calif.merge_cells("A5:A7")
    ws_calif["A5"] = "Matrícula"
    ws_calif["A5"].font = bold_font
    ws_calif["A5"].alignment = Alignment(horizontal="center", vertical="center")
    ws_calif["A5"].border = thin_border
    
    # B5:B7 para Nombre
    ws_calif.merge_cells("B5:B7")
    ws_calif["B5"] = "Nombre"
    ws_calif["B5"].font = bold_font
    ws_calif["B5"].alignment = Alignment(horizontal="center", vertical="center")
    ws_calif["B5"].border = thin_border
    
    current_col = 3
    for p in ponderaciones:
        acts = p.get("actividades", [])
        span = len(acts) + 1 # Número de actividades más la columna de total tareas/subtotal
        end_col = current_col + span - 1
        
        start_letter = get_column_letter(current_col)
        end_letter = get_column_letter(end_col)
        
        # Fila 5: Porcentaje de la categoría (ej. "10%")
        ws_calif.merge_cells(f"{start_letter}5:{end_letter}5")
        cell_pct = ws_calif[f"{start_letter}5"]
        cell_pct.value = f"{int(float(p.get('porcentaje', 0)))}%"
        cell_pct.font = bold_font
        cell_pct.alignment = Alignment(horizontal="center", vertical="center")
        cell_pct.border = thin_border
        
        # Fila 6: Nombre descriptivo de la categoría
        ws_calif.merge_cells(f"{start_letter}6:{end_letter}6")
        cell_desc = ws_calif[f"{start_letter}6"]
        cell_desc.value = p.get("nombre_categoria", "Categoría")
        cell_desc.font = Font(name="Calibri", size=9, bold=True)
        cell_desc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_desc.border = thin_border
        
        # Fila 7: Actividades específicas y subtotal de la categoría
        for act in acts:
            cell_act = ws_calif.cell(row=7, column=current_col)
            cell_act.value = act.get("nombre", "Actividad")
            cell_act.font = Font(name="Calibri", size=8, bold=True)
            # Rotación del texto a 90 grados para encajar nombres largos
            cell_act.alignment = Alignment(text_rotation=90, wrap_text=True, horizontal="center", vertical="center")
            cell_act.border = thin_border
            current_col += 1
            
        # Columna de subtotal intermedio
        cell_sub = ws_calif.cell(row=7, column=current_col)
        alias = p.get("nombre_categoria", "TAREAS").split(" ")[0].upper()
        cell_sub.value = f"TOTAL {alias}"
        cell_sub.font = Font(name="Calibri", size=9, bold=True)
        cell_sub.alignment = Alignment(text_rotation=90, wrap_text=True, horizontal="center", vertical="center")
        cell_sub.fill = fill_yellow
        cell_sub.border = thin_border
        current_col += 1
        
    # Calificación Final combinada (Columna final, filas 5 a 7)
    ws_calif.merge_cells(f"{last_col_letter}5:{last_col_letter}7")
    cell_final = ws_calif[f"{last_col_letter}5"]
    cell_final.value = "Calificación\nFinal"
    cell_final.font = bold_font
    cell_final.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_final.border = thin_border
    
    # Asignar bordes para todas las celdas de encabezados de tabla (filas 5 a 7)
    for r in range(5, 8):
        for c in range(1, last_col_idx + 1):
            ws_calif.cell(row=r, column=c).border = thin_border
            
    # Ajustar altura de fila 7 para texto vertical
    ws_calif.row_dimensions[7].height = 130
    
    # 3. Escribir Datos de los Alumnos
    row_idx = 8
    alumnos = datos_calificaciones.get("alumnos", [])
    
    for al in alumnos:
        ws_calif.cell(row=row_idx, column=1, value=al.get("matricula", "N/A")).alignment = Alignment(horizontal="center")
        ws_calif.cell(row=row_idx, column=2, value=al.get("alumno_nombre", "Desconocido")).alignment = Alignment(horizontal="left")
        
        current_col = 3
        alumno_calificaciones = al.get("calificaciones", {})
        
        for p in ponderaciones:
            acts = p.get("actividades", [])
            total_cat = 0.0
            valid_acts_count = 0
            
            for act in acts:
                # Recuperar calificación individual
                val_raw = alumno_calificaciones.get(act.get("id"), 0.0)
                # Escalar de 0-10 a 0-1
                val_scaled = round(float(val_raw) / 10.0, 2)
                
                cell_val = ws_calif.cell(row=row_idx, column=current_col, value=val_scaled)
                cell_val.alignment = Alignment(horizontal="center")
                current_col += 1
                
                total_cat += val_scaled
                valid_acts_count += 1
                
            # Calcular subtotal intermedio
            cat_average = round(total_cat / valid_acts_count, 2) if valid_acts_count > 0 else 0.0
            cell_sub_val = ws_calif.cell(row=row_idx, column=current_col, value=cat_average)
            cell_sub_val.alignment = Alignment(horizontal="center")
            cell_sub_val.font = bold_font
            cell_sub_val.fill = fill_yellow
            current_col += 1
            
        # Calificación Final (en escala 0-10)
        final_grade = al.get("promedio_real", 0.0)
        cell_final_val = ws_calif.cell(row=row_idx, column=current_col, value=final_grade)
        cell_final_val.alignment = Alignment(horizontal="center")
        cell_final_val.font = bold_font
        
        # Aplicar fuente regular y bordes a toda la fila
        for col_i in range(1, last_col_idx + 1):
            c_cell = ws_calif.cell(row=row_idx, column=col_i)
            if col_i != 2 and col_i != last_col_idx:
                c_cell.font = regular_font
            c_cell.border = thin_border
            
        row_idx += 1
        
    # Ajustar anchos
    ws_calif.column_dimensions['A'].width = 15
    ws_calif.column_dimensions['B'].width = 38
    for col_i in range(3, last_col_idx):
        ws_calif.column_dimensions[get_column_letter(col_i)].width = 7
    ws_calif.column_dimensions[last_col_letter].width = 14
    
    # -------------------------------------------------------------------------
    # HOJA 2: ASISTENCIAS
    # -------------------------------------------------------------------------
    ws_asist = wb.create_sheet(title="Asistencias")
    ws_asist.views.sheetView[0].showGridLines = True
    
    # 1. Cabecera (Filas 1 a 5)
    # Recolectar fechas únicas de todas las asistencias registradas
    fechas_unicas = set()
    for s_asist in datos_asistencias:
        for item in s_asist.get('asistencias', []):
            fechas_unicas.add(item['fecha'])
    fechas_ordenadas = sorted(list(fechas_unicas))
    
    num_fechas = len(fechas_ordenadas)
    asist_last_col_idx = 2 + num_fechas
    asist_last_col_letter = get_column_letter(asist_last_col_idx)
    
    # Escribir cabecera combinada para Asistencias
    cabecera_asist_textos = [
        materia_nombre,
        "BENEMÉRITA UNIVERSIDAD AUTÓNOMA DE PUEBLA",
        "FACULTAD DE CIENCIAS DE LA COMPUTACIÓN",
        docente_str,
        periodo_str
    ]
    
    for idx, text in enumerate(cabecera_asist_textos, start=1):
        range_str = f"A{idx}:{asist_last_col_letter}{idx}"
        ws_asist.merge_cells(range_str)
        cell = ws_asist[f"A{idx}"]
        cell.value = text
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = fill_blue_light
        
    for r in range(1, 6):
        for c in range(1, asist_last_col_idx + 1):
            cell = ws_asist.cell(row=r, column=c)
            cell.fill = fill_blue_light
            cell.border = thin_border
            
    # 2. Encabezados de la Tabla (Fila 6)
    ws_asist.cell(row=6, column=1, value="Matrícula").font = bold_font
    ws_asist.cell(row=6, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws_asist.cell(row=6, column=2, value="Nombre").font = bold_font
    ws_asist.cell(row=6, column=2).alignment = Alignment(horizontal="center", vertical="center")
    
    for idx, fecha in enumerate(fechas_ordenadas, start=3):
        fecha_formateada = _format_fecha_buap(fecha)
        cell_f = ws_asist.cell(row=6, column=idx, value=fecha_formateada)
        cell_f.font = Font(name="Calibri", size=9, bold=True)
        cell_f.alignment = Alignment(text_rotation=90, horizontal="center", vertical="center")
        
    # Bordes de fila 6
    for c in range(1, asist_last_col_idx + 1):
        ws_asist.cell(row=6, column=c).border = thin_border
        
    ws_asist.row_dimensions[6].height = 65
    
    # 3. Escribir Datos de Asistencias (Fila 7 en adelante)
    row_idx = 7
    # Mapeo rápido de asistencias por alumno_id para optimizar búsqueda
    asist_map = {}
    for sa in datos_asistencias:
        al_id = str(sa.get("alumno_id"))
        recs = {}
        for item in sa.get("asistencias", []):
            recs[item["fecha"]] = item["estado"]
        asist_map[al_id] = recs
        
    # Obtener alumnos del concentrado para tener consistencia e ID
    for al in alumnos:
        al_id = str(al.get("alumno_id"))
        matricula = al.get("matricula", "N/A")
        nombre = al.get("alumno_nombre", "Desconocido")
        
        ws_asist.cell(row=row_idx, column=1, value=matricula).alignment = Alignment(horizontal="center")
        ws_asist.cell(row=row_idx, column=2, value=nombre).alignment = Alignment(horizontal="left")
        
        student_records = asist_map.get(al_id, {})
        
        for idx, fecha in enumerate(fechas_ordenadas, start=3):
            estado = student_records.get(fecha, "ausente").lower()
            
            # Regla de Negocio: 1 = presente, 0.5 = retardo, 0 = falta/ausente
            if estado == "presente":
                val = 1
            elif estado == "retardo":
                val = 0.5
            else:
                val = 0
                
            cell_val = ws_asist.cell(row=row_idx, column=idx, value=val)
            cell_val.alignment = Alignment(horizontal="center")
            
        # Estilos y bordes a toda la fila
        for col_i in range(1, asist_last_col_idx + 1):
            c_cell = ws_asist.cell(row=row_idx, column=col_i)
            c_cell.font = regular_font if col_i == 2 else bold_font if col_i > 2 else regular_font
            c_cell.border = thin_border
            
        row_idx += 1
        
    # Ajustar anchos en Asistencias
    ws_asist.column_dimensions['A'].width = 15
    ws_asist.column_dimensions['B'].width = 38
    for col_i in range(3, asist_last_col_idx + 1):
        ws_asist.column_dimensions[get_column_letter(col_i)].width = 9
        
    # 4. Guardar y retornar
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()